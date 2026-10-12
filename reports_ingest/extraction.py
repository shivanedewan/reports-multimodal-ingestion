from __future__ import annotations

import bz2
import csv
import email
import gzip
import io
import json
import lzma
import math
import mimetypes
import os
import shutil
import stat
import subprocess
import tarfile
import tempfile
import wave
import zipfile
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from .common import GlobalFailure, IngestError, digest, extension, file_hash, normalized_text, usable_text

TEXT = set("txt text log md rst csv tsv json jsonl xml html htm".split())
WORD = set("doc docx docm dot dotx dotm rtf odt ott fodt wpd sxw".split())
SHEETS = set("xls xlsx xlsm xlsb ods fods".split())
SLIDES = set("ppt pptx pptm pps ppsx odp fodp".split())
IMAGES = set("jpg jpeg png tif tiff bmp webp gif heic heif jp2".split())
AUDIO = set("wav mp3 flac m4a aac ogg opus wma aif aiff amr au".split())
VIDEO = set("mp4 mkv avi mov wmv mpeg mpg webm m4v 3gp".split())
ARCHIVES = set("zip tar gz tgz bz2 xz tar.gz tar.bz2 tar.xz 7z rar".split())
LEGACY = set("eml msg mbox pst epub mhtml mht".split())


@dataclass
class Unit:
    modality: str
    fields: dict
    text: str | None = None
    media: bytes | None = None
    mime: str | None = None


@dataclass
class Notice:
    code: str
    reason: str
    fields: dict = field(default_factory=dict)
    disposition: str = "unit_skipped"
    affects_coverage: bool = True
    stage: str = "extract"


def registry(config):
    return {"text": sorted(TEXT), "word": sorted(WORD), "spreadsheets": sorted(SHEETS),
        "presentations": sorted(SLIDES), "images": sorted(IMAGES), "audio": sorted(AUDIO),
        "archives": sorted(ARCHIVES), "other": sorted(LEGACY), "video": sorted(VIDEO),
        "dependencies": {"ffmpeg": shutil.which(config.extract["ffmpeg"]), "ffprobe": shutil.which(config.extract["ffprobe"]),
            "tika": bool(config.extract["tika_url"]), "libreoffice": config.extract["libreoffice"], "ocr": config.extract["tesseract"]},
        "notes": ["Legacy Office/MSG/mail stores use Tika or configured conversion; EPUB has section text extraction",
            "7z safe expansion is unavailable: ES Text fallback or logged skip; RAR requires optional rarfile/backend", "Video is opt-in"]}


def media_type(ext):
    return "image" if ext in IMAGES else "audio" if ext in AUDIO else "video" if ext in VIDEO else "archive" if ext in ARCHIVES else "document"


def detect(path: Path, declared: str):
    with path.open("rb") as stream:
        header = stream.read(32)
    if header.startswith(b"%PDF"):
        return "pdf", "application/pdf"
    if header.startswith(b"PK\x03\x04"):
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            if "word/document.xml" in names:
                return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if "xl/workbook.xml" in names:
                return "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if "ppt/presentation.xml" in names:
                return "pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            if "mimetype" in names:
                info = package.getinfo("mimetype")
                if info.file_size <= 256:
                    mime = package.read(info).decode("ascii", "replace")
                    odf = {"application/vnd.oasis.opendocument.text": "odt", "application/vnd.oasis.opendocument.spreadsheet": "ods", "application/vnd.oasis.opendocument.presentation": "odp"}
                    if mime in odf:
                        return odf[mime], mime
        return "zip", "application/zip"
    signatures = [(b"\x89PNG", "png"), (b"\xff\xd8\xff", "jpg"), (b"GIF8", "gif"),
        (b"II*\x00", "tiff"), (b"MM\x00*", "tiff"), (b"BM", "bmp"), (b"fLaC", "flac"), (b"ID3", "mp3"),
        (b"7z\xbc\xaf\x27\x1c", "7z"), (b"Rar!", "rar")]
    for signature, ext in signatures:
        if header.startswith(signature):
            return ext, mimetypes.guess_type("x." + ext)[0] or "application/octet-stream"
    if header.startswith(b"RIFF") and header[8:12] == b"WAVE":
        return "wav", "audio/wav"
    if header.startswith(b"RIFF") and header[8:12] == b"WEBP":
        return "webp", "image/webp"
    if declared not in TEXT | WORD | SHEETS | SLIDES | IMAGES | AUDIO | VIDEO | ARCHIVES | LEGACY | {"pdf"}:
        try:
            sample = header.decode("utf-8")
            if usable_text(sample):
                return "txt", "text/plain"
        except UnicodeDecodeError:
            pass
    return declared, mimetypes.guess_type("x." + declared)[0] or "application/octet-stream"


def command(args: list[str], timeout: float) -> bytes:
    try:
        result = subprocess.run(args, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            timeout=timeout, check=False, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except FileNotFoundError:
        raise IngestError("dependency_missing", f"Executable unavailable: {Path(args[0]).name}") from None
    except subprocess.TimeoutExpired:
        raise IngestError("parser_timeout", f"{Path(args[0]).name} exceeded {timeout}s") from None
    if result.returncode:
        # Converter stderr may contain extracted text; avoid copying it to logs.
        raise IngestError("decoder_failed", f"{Path(args[0]).name} returned exit status {result.returncode}")
    return result.stdout


class Extractor:
    def __init__(self, config):
        self.config, self.settings = config, config.extract

    def path(self, original):
        text = str(original or "")
        for old, new in self.settings["path_prefixes"].items():
            if text.startswith(old):
                text = new + text[len(old):]
                if os.name != "nt":
                    text = text.replace("\\", "/")
                break
        return Path(text) if text else None

    def record(self, hit: dict):
        source = hit["_source"]
        original = source.get("OriginalName") or str(source.get("SystemPath") or "")
        ext = extension(str(source.get("FileExtension") or Path(original).name))
        path = self.path(source.get("SystemPath"))
        text = normalized_text(source.get("Text"))
        base = {"SourceMediaType": media_type(ext), "FileAvailability": "unknown", "UsedTextFallback": False,
            "PreviewAvailable": False, "ContentSource": "file", "RepresentationKind": "native", "ArchiveChain": []}
        try:
            if path is None:
                raise FileNotFoundError("Source has no SystemPath")
            info = path.stat()
            if not path.is_file():
                raise FileNotFoundError("SystemPath does not identify a regular file")
            base.update(FileAvailability="present", FileExists=True, PreviewAvailable=True,
                FileSizeBytes=info.st_size, FileModifiedAt=__import__("datetime").datetime.fromtimestamp(info.st_mtime, __import__("datetime").timezone.utc).isoformat())
            with path.open("rb") as probe:
                probe.read(1)
        except (FileNotFoundError, PermissionError, OSError) as error:
            missing = isinstance(error, FileNotFoundError)
            base.update(FileAvailability="missing" if missing else "unreadable", FileExists=False if missing else None)
            code = "missing_file" if missing else "file_access_failed"
            if usable_text(text):
                yield Notice(code, "Original unavailable; embedding usable ES Text", base, "fallback_indexed", False)
                yield self.text_unit(text, {**base, "UsedTextFallback": True, "ContentSource": "es_text", "RepresentationKind": "es_text_fallback"})
            else:
                yield Notice(code + "_no_text", "Original unavailable and no usable ES Text", base, "source_skipped")
            return
        # Ordinary prose uses ES Text before touching complex parsers. Signature detection avoids
        # losing an incorrectly named media file behind a text-only fast path.
        emitted = False
        try:
            if ext in ARCHIVES and path.stat().st_size > self.settings.get("max_archive_bytes", 2147483648):
                raise IngestError("archive_input_size_limit", "Archive input exceeds configured byte budget")
            actual_ext, mime = detect(path, ext)
            base.update(DetectedMimeType=mime, SourceMediaType=media_type(actual_ext))
            if actual_ext != ext:
                yield Notice("extension_mismatch", f"Declared extension {ext!r}; signature suggests {actual_ext!r}", base, "processed_by_signature", False)
            ext = actual_ext
            limit = self.settings["max_media_bytes"] if ext in AUDIO | VIDEO else self.settings["max_document_bytes"]
            if ext in ARCHIVES and path.stat().st_size > self.settings.get("max_archive_bytes", 2147483648):
                raise IngestError("archive_input_size_limit", "Detected archive exceeds configured input byte budget")
            if ext not in ARCHIVES and path.stat().st_size > limit:
                raise IngestError("file_size_limit", "File exceeds configured input byte limit")
            # XLS/PPT/PDF use anchors/structure from the file, with ES fallback.
            if usable_text(text) and ext in TEXT | WORD:
                yield self.text_unit(text, {**base, "ContentSource": "es_text", "RepresentationKind": "prose"})
                return
            for item in self.file(path, ext, base, depth=0, budget={"bytes": 0, "members": 0}):
                if isinstance(item, Unit):
                    emitted = True
                yield item
            if not emitted:
                if usable_text(text):
                    yield Notice("no_extracted_units", "File yielded no usable units; embedding ES Text fallback", base, "fallback_indexed", False)
                    yield self.text_unit(text, {**base, "UsedTextFallback": True, "ContentSource": "es_text", "RepresentationKind": "es_text_fallback"})
                else:
                    yield Notice("empty_content", "File yielded no usable text/media", base, "source_skipped")
            elif usable_text(text) and ext in IMAGES | AUDIO | VIDEO:
                yield self.text_unit(text, {**base, "ContentSource": "es_text", "RepresentationKind": "associated_text"})
        except GlobalFailure:
            raise
        except Exception as error:
            code = getattr(error, "code", "file_parse_failed")
            if emitted:
                yield Notice(code, "File partially processed; no duplicate aggregate ES Text fallback", base, "source_partial")
            elif usable_text(text):
                yield Notice(code, "File parsing failed; embedding usable ES Text fallback", base, "fallback_indexed", False)
                yield self.text_unit(text, {**base, "UsedTextFallback": True, "ContentSource": "es_text", "RepresentationKind": "es_text_fallback"})
            else:
                yield Notice(code, str(error) if isinstance(error, IngestError) else type(error).__name__, base, "source_skipped")

    def text_unit(self, text, fields):
        return Unit("text", {"ExtractionMethod": "es_text" if fields.get("ContentSource") == "es_text" else "native_text",
            "TextQualityStatus": "usable", **fields}, text=normalized_text(text))

    def file(self, path: Path, ext: str, fields: dict, depth: int, budget: dict):
        if ext in ARCHIVES:
            yield from self.archive(path, ext, fields, depth, budget)
        elif ext == "pdf":
            yield from self.pdf(path, fields)
        elif ext in IMAGES:
            yield from self.image(path, fields)
        elif ext in AUDIO:
            yield from self.audio(path, fields)
        elif ext in VIDEO:
            if not self.settings["enable_video"]:
                raise IngestError("video_disabled", "Enable video adapter to process this source")
            yield from self.video(path, fields)
        elif ext in {"xlsx", "xlsm"}:
            yield from self.sheet(path, fields)
        elif ext in {"odt", "ott", "ods", "odp", "epub", "docx", "pptx"}:
            yield from self.packaged(path, ext, fields)
        elif ext in TEXT:
            yield from self.plain(path, ext, fields)
        elif ext == "eml":
            with path.open("rb") as stream:
                message = email.message_from_binary_file(stream)
            pieces = [f"Subject: {message.get('Subject', '')}"]
            for part in message.walk():
                if part.get_content_type() in {"text/plain", "text/html"} and part.get_content_disposition() != "attachment":
                    raw = part.get_payload(decode=True) or b""
                    body = raw.decode(part.get_content_charset() or "utf-8", "replace")
                    if part.get_content_type() == "text/html":
                        from bs4 import BeautifulSoup
                        body = BeautifulSoup(body, "html.parser").get_text("\n")
                    pieces.append(body)
            yield self.text_unit("\n\n".join(pieces), fields)
            # Embedded MIME attachments need separate locators; do not pretend they were indexed.
            if any(p.get_content_disposition() == "attachment" for p in message.walk()):
                yield Notice("mail_attachments_not_expanded", "EML body indexed; MIME attachments require separate archive adapter", fields)
        else:
            yield from self.legacy(path, ext, fields)

    def plain(self, path, ext, fields):
        from charset_normalizer import from_bytes
        raw = path.read_bytes()
        text = str(from_bytes(raw).best() or "")
        if ext in {"html", "htm", "xml"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(text, "html.parser")
            for node in soup(["script", "style"]):
                node.decompose()
            text = soup.get_text("\n")
        if ext in {"csv", "tsv"}:
            rows = csv.reader(io.StringIO(text), delimiter="\t" if ext == "tsv" else ",")
            header = next(rows, [])
            for number, row in enumerate(rows, 2):
                body = " | ".join(f"{header[i] if i < len(header) else f'Column {i+1}'}: {v}" for i, v in enumerate(row))
                if usable_text(body):
                    yield self.text_unit(body, {**fields, "RowStart": number, "RowEnd": number, "RepresentationKind": "table_row"})
            return
        if usable_text(text):
            yield self.text_unit(text, fields)

    def packaged(self, path, ext, fields):
        # Stream bounded ZIP XML parts without executing macros or expanding arbitrary binaries.
        with zipfile.ZipFile(path) as package:
            if ext == "docx":
                names = [n for n in package.namelist() if n == "word/document.xml" or n.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")) and n.endswith(".xml")]
            elif ext == "pptx":
                names = sorted((n for n in package.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")), key=lambda n: int(Path(n).stem.replace("slide", "")))
            elif ext == "epub":
                # Reading order isn't guaranteed here; use section locator, no false page numbers.
                names = [n for n in package.namelist() if n.endswith((".xhtml", ".html", ".htm"))]
            else:
                names = ["content.xml"]
            if len(names) > self.settings["archive_members"]:
                raise IngestError("package_part_limit", "Document package exceeds part limit")
            total = 0
            for number, name in enumerate(names, 1):
                info = package.getinfo(name)
                total += info.file_size
                if total > self.settings["archive_expanded_bytes"] or info.file_size > self.settings["max_document_bytes"]:
                    raise IngestError("package_size_limit", "Document XML exceeds expansion limit")
                data = package.read(info)
                if ext == "epub":
                    from bs4 import BeautifulSoup
                    text = BeautifulSoup(data, "html.parser").get_text("\n")
                else:
                    root = ET.fromstring(data)
                    text = "\n".join(t for t in root.itertext() if t.strip())
                if usable_text(text):
                    yield self.text_unit(text, {**fields, "SectionPath": name,
                        **({"SlideNumber": number} if ext == "pptx" else {})})
            if ext in {"docx", "pptx", "odt", "ods", "odp"} and any(n.startswith(("word/media/", "ppt/media/", "Pictures/")) for n in package.namelist()):
                yield Notice("office_visuals_not_expanded", "Text indexed; embedded Office images are not represented in this pass", fields)

    def sheet(self, path, fields):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                header = None
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = ["" if v is None else str(v) for v in row]
                    if not any(values):
                        continue
                    if header is None:
                        header = values
                        text = "Headers: " + " | ".join(values)
                    else:
                        text = " | ".join(f"{header[i] if i < len(header) and header[i] else f'Column {i+1}'}: {v}" for i, v in enumerate(values))
                    yield self.text_unit(text, {**fields, "SheetName": sheet.title, "RowStart": number, "RowEnd": number,
                        "RepresentationKind": "table_row", "ExtractionMethod": "openpyxl"})
        finally:
            book.close()

    def legacy(self, path, ext, fields):
        if self.settings["tika_url"]:
            import httpx
            with path.open("rb") as stream:
                with httpx.Client(timeout=self.settings["command_timeout_seconds"], trust_env=False) as client:
                    response = client.put(self.settings["tika_url"].rstrip("/") + "/tika", content=stream,
                        headers={"Accept": "text/plain", "X-Tika-OCRskipOcr": "true"})
            if response.status_code >= 400:
                raise IngestError("tika_failed", f"Tika returned HTTP {response.status_code}")
            if usable_text(response.text):
                yield self.text_unit(response.text, {**fields, "ExtractionMethod": "tika"})
            return
        office = self.settings["libreoffice"]
        if office and ext in WORD | SHEETS | SLIDES:
            with tempfile.TemporaryDirectory(dir=self.config.path(self.config.runtime["temp_path"])) as work:
                work = Path(work)
                profile = work / "profile"
                command([office, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "txt:Text", "--outdir", str(work), str(path)], self.settings["command_timeout_seconds"])
                output = work / (path.stem + ".txt")
                if not output.is_file():
                    raise IngestError("conversion_empty", "LibreOffice did not produce text output")
                yield from self.plain(output, "txt", {**fields, "ExtractionMethod": "libreoffice"})
            return
        raise IngestError("unsupported_or_dependency_missing", f"No configured parser for .{ext}; configure Tika/converter or use ES Text")

    def image(self, path, fields):
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = self.settings["max_image_pixels"]
        with Image.open(path) as image:
            frames = getattr(image, "n_frames", 1)
            maximum = self.settings["gif_frames"] if image.format == "GIF" else self.settings["max_pdf_pages"]
            selected = sorted({int(i * (frames - 1) / max(1, min(frames, maximum) - 1)) for i in range(min(frames, maximum))})
            if len(selected) < frames:
                yield Notice("image_frame_limit", f"Processing {len(selected)} of {frames} frames/pages under configured policy", fields)
            for index in selected:
                try:
                    image.seek(index)
                    original_size = image.size
                    if image.width * image.height > self.settings["max_image_pixels"]:
                        raise IngestError("image_pixel_limit", "Image exceeds configured decoded pixel limit")
                    visual = ImageOps.exif_transpose(image).convert("RGB")
                    visual.thumbnail((self.settings["max_image_side"], self.settings["max_image_side"]))
                    buffer = io.BytesIO()
                    visual.save(buffer, format="PNG")
                    extra = {"ImageWidth": original_size[0], "ImageHeight": original_size[1], "FrameNumber": index,
                        "RepresentationKind": "image_frame" if frames > 1 else "image", "ExtractionMethod": "pillow"}
                    if image.format == "TIFF":
                        extra.update(PageStart=index + 1, PageEnd=index + 1)
                    yield Unit("image", {**fields, **extra}, media=buffer.getvalue(), mime="image/png")
                    if self.settings["tesseract"]:
                        with tempfile.TemporaryDirectory(dir=self.config.path(self.config.runtime["temp_path"])) as work:
                            png = Path(work) / "ocr.png"
                            visual.save(png)
                            text = command([self.settings["tesseract"], str(png), "stdout", "-l", self.settings["ocr_language"]], self.settings["command_timeout_seconds"]).decode("utf-8", "replace")
                            if usable_text(text):
                                yield self.text_unit(text, {**fields, **extra, "RepresentationKind": "ocr", "ExtractionMethod": "tesseract"})
                except Exception as error:
                    yield Notice(getattr(error, "code", "image_frame_failed"), str(error) if isinstance(error, IngestError) else type(error).__name__, {**fields, "FrameNumber": index})

    def pdf(self, path, fields):
        try:
            import pymupdf
        except ImportError:
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise IngestError("encrypted_pdf", "PDF requires a password")
            pages = len(reader.pages)
            for number, page in enumerate(reader.pages[:self.settings["max_pdf_pages"]], 1):
                try:
                    text = page.extract_text() or ""
                    anchor = {**fields, "PageStart": number, "PageEnd": number, "ExtractionMethod": "pypdf"}
                    if usable_text(text):
                        yield self.text_unit(text, anchor)
                    else:
                        yield Notice("pdf_renderer_missing", "Scanned page requires optional PyMuPDF renderer", anchor)
                except Exception as error:
                    yield Notice("pdf_page_failed", type(error).__name__, {**fields, "PageStart": number})
            if pages > self.settings["max_pdf_pages"]:
                yield Notice("pdf_page_limit", "PDF exceeds configured page limit", fields)
            return
        with pymupdf.open(path) as document:
            if document.needs_pass and not document.authenticate(""):
                raise IngestError("encrypted_pdf", "PDF requires a password")
            for index in range(min(len(document), self.settings["max_pdf_pages"])):
                anchor = {**fields, "PageStart": index + 1, "PageEnd": index + 1, "ExtractionMethod": "pymupdf"}
                try:
                    page = document[index]
                    text = page.get_text()
                    has_text = usable_text(text)
                    if has_text:
                        yield self.text_unit(text, {**anchor, "RepresentationKind": "pdf_text"})
                    mode = self.settings["pdf_visual_mode"]
                    if mode == "all" or not has_text or mode == "mixed" and page.get_images():
                        scale = min(2.0, self.settings["max_image_side"] / max(page.rect.width, page.rect.height))
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                        png = pix.tobytes("png")
                        yield Unit("image", {**anchor, "RepresentationKind": "pdf_page", "ImageWidth": pix.width, "ImageHeight": pix.height}, media=png, mime="image/png")
                        if not has_text and self.settings["tesseract"]:
                            with tempfile.TemporaryDirectory(dir=self.config.path(self.config.runtime["temp_path"])) as work:
                                target = Path(work) / "page.png"
                                target.write_bytes(png)
                                ocr = command([self.settings["tesseract"], str(target), "stdout", "-l", self.settings["ocr_language"]], self.settings["command_timeout_seconds"]).decode("utf-8", "replace")
                                if usable_text(ocr):
                                    yield self.text_unit(ocr, {**anchor, "RepresentationKind": "ocr", "ExtractionMethod": "tesseract"})
                        elif not has_text:
                            yield Notice("ocr_not_configured", "Visual page indexed; no OCR text representation", anchor, "visual_only", False)
                    elif page.get_images():
                        yield Notice("pdf_visual_pass_disabled", "Text indexed; page images require pdf_visual_mode=mixed/all", anchor)
                except Exception as error:
                    yield Notice(getattr(error, "code", "pdf_page_failed"), str(error) if isinstance(error, IngestError) else type(error).__name__, anchor)
            if len(document) > self.settings["max_pdf_pages"]:
                yield Notice("pdf_page_limit", "PDF exceeds configured page limit", fields)

    def probe(self, path):
        raw = command([self.settings["ffprobe"], "-v", "error", "-show_entries", "format=duration:stream=codec_type,duration", "-of", "json", str(path)], self.settings["command_timeout_seconds"])
        info = json.loads(raw)
        try:
            duration = float(info.get("format", {}).get("duration") or max(float(s.get("duration", 0)) for s in info.get("streams", [])))
        except (ValueError, TypeError):
            raise IngestError("media_duration_unknown", "Cannot determine finite media duration") from None
        if not math.isfinite(duration) or duration <= 0:
            raise IngestError("media_duration_unknown", "Cannot determine finite media duration")
        return duration, {s.get("codec_type") for s in info.get("streams", [])}

    def audio(self, path, fields):
        duration, streams = self.probe(path)
        if "audio" not in streams:
            raise IngestError("no_audio_stream", "Container has no audio stream")
        window = self.settings["audio_window_seconds"]
        stride = window - self.settings["audio_overlap_seconds"]
        start = 0.0
        while start < duration:
            end = min(duration, start + window)
            anchor = {**fields, "StartSeconds": start, "EndSeconds": end, "MediaDurationSeconds": duration,
                "RepresentationKind": "audio_segment", "ExtractionMethod": "ffmpeg"}
            try:
                data = command([self.settings["ffmpeg"], "-v", "error", "-nostdin", "-ss", str(start), "-i", str(path),
                    "-t", str(end - start), "-vn", "-ac", "1", "-ar", "16000", "-f", "s16le", "pipe:1"], self.settings["command_timeout_seconds"])
                if not data:
                    raise IngestError("empty_audio_segment", "Decoder returned no samples")
                actual_seconds = len(data) / 32000
                anchor["EndSeconds"] = min(end, start + actual_seconds)
                output = io.BytesIO()
                with wave.open(output, "wb") as wav:
                    wav.setnchannels(1); wav.setsampwidth(2); wav.setframerate(16000); wav.writeframes(data)
                yield Unit("audio", anchor, media=output.getvalue(), mime="audio/wav")
            except IngestError as error:
                yield Notice(error.code, str(error), anchor)
            if end == duration:
                break
            start += stride

    def video(self, path, fields):
        duration, streams = self.probe(path)
        if "audio" in streams:
            yield from self.audio(path, fields)
        if "video" not in streams:
            return
        interval = self.settings["video_frame_interval_seconds"]
        for index in range(math.ceil(duration / interval)):
            start = index * interval
            anchor = {**fields, "StartSeconds": start, "EndSeconds": min(duration, start + interval),
                "MediaDurationSeconds": duration, "FrameNumber": index, "RepresentationKind": "video_sampled_frame", "ExtractionMethod": "ffmpeg"}
            try:
                data = command([self.settings["ffmpeg"], "-v", "error", "-nostdin", "-ss", str(start), "-i", str(path),
                    "-frames:v", "1", "-vf", f"scale={self.settings['max_image_side']}:{self.settings['max_image_side']}:force_original_aspect_ratio=decrease", "-f", "image2pipe", "-vcodec", "png", "pipe:1"], self.settings["command_timeout_seconds"])
                if not data:
                    raise IngestError("empty_video_frame", "Decoder produced no visual frame")
                yield Unit("image", anchor, media=data, mime="image/png")
            except IngestError as error:
                yield Notice(error.code, str(error), anchor)
        yield Notice("video_sampling_policy", "Indexed sampled frames and audio; native temporal-video embeddings are not generated", fields, "sampled_video", False)

    def archive(self, path, ext, fields, depth, budget):
        if depth >= self.settings["archive_depth"]:
            raise IngestError("archive_depth_limit", "Nested archive depth exceeds configured limit")
        temp = self.config.path(self.config.runtime["temp_path"])
        with tempfile.TemporaryDirectory(dir=temp) as work:
            if ext == "zip":
                with zipfile.ZipFile(path) as archive:
                    for ordinal, member in enumerate(archive.infolist()):
                        if budget["members"] >= self.settings["archive_members"] or budget["bytes"] >= self.settings["archive_expanded_bytes"]:
                            yield Notice("archive_budget_exhausted", "Remaining archive entries excluded under resource policy", fields, "remaining_members_skipped")
                            break
                        if member.is_dir():
                            continue
                        if member.flag_bits & 1:
                            yield Notice("encrypted_archive_member", "ZIP member requires a password", {**fields, "MemberPath": member.filename}, "member_skipped")
                            budget["members"] += 1
                            continue
                        mode = (member.external_attr >> 16) & 0xFFFF
                        yield from self.member(path, member.filename, ordinal, member.file_size, member.compress_size,
                            lambda m=member: archive.open(m), bool(stat.S_ISLNK(mode)), fields, depth, budget, Path(work))
            elif ext in {"tar", "tgz", "tar.gz", "tar.bz2", "tar.xz"}:
                with tarfile.open(path, "r|*") as archive:
                    for ordinal, member in enumerate(archive):
                        if budget["members"] >= self.settings["archive_members"] or budget["bytes"] >= self.settings["archive_expanded_bytes"]:
                            yield Notice("archive_budget_exhausted", "Remaining archive entries excluded under resource policy", fields, "remaining_members_skipped")
                            break
                        if member.isdir():
                            continue
                        yield from self.member(path, member.name, ordinal, member.size, None,
                            lambda m=member: archive.extractfile(m), not member.isfile(), fields, depth, budget, Path(work))
            elif ext == "rar":
                try:
                    import rarfile
                except ImportError:
                    raise IngestError("dependency_missing", "Install rarfile and its extraction backend") from None
                with rarfile.RarFile(path) as archive:
                    for ordinal, member in enumerate(archive.infolist()):
                        if member.isdir():
                            continue
                        yield from self.member(path, member.filename, ordinal, member.file_size, member.compress_size,
                            lambda m=member: archive.open(m), member.is_symlink(), fields, depth, budget, Path(work))
            elif ext == "7z":
                # py7zr's extractall does not enforce actual expanded-byte budgets while streaming.
                # Refuse unsafe bulk expansion; Tika aggregate text or ES Text can rescue the source.
                raise IngestError("safe_7z_adapter_unavailable", "7z expansion requires a bounded streaming adapter; use ZIP/TAR/RAR or ES Text fallback")
            else:
                opener = {"gz": gzip.open, "bz2": bz2.open, "xz": lzma.open}.get(ext)
                if not opener:
                    raise IngestError("unsupported_archive", f"No archive adapter for {ext}")
                name = path.name.rsplit(".", 1)[0]
                yield from self.member(path, name, 0, None, None, lambda: opener(path, "rb"), False, fields, depth, budget, Path(work))

    def member(self, container, name, ordinal, size, compressed, opener, unsafe, fields, depth, budget, work):
        member_fields = {**fields, "ArchiveChain": [*fields.get("ArchiveChain", []), {"name": container.name, "member": name, "ordinal": ordinal}],
            "MemberPath": name, "MemberOrdinal": ordinal, "ArchiveDepth": depth + 1, "MemberId": digest([fields.get("ArchiveChain", []), name, ordinal]),
            "ContentSource": "archive_member", "PreviewAvailable": True}
        budget["members"] += 1
        try:
            normalized = name.replace("\\", "/")
            parts = PurePosixPath(normalized)
            if unsafe or parts.is_absolute() or ".." in parts.parts or ":" in normalized or "\x00" in normalized:
                raise IngestError("unsafe_archive_member", "Archive member path/link is unsafe")
            if budget["members"] > self.settings["archive_members"]:
                raise IngestError("archive_member_limit", "Archive exceeds member count limit")
            ext = extension(Path(normalized).name)
            limit = self.settings["max_media_bytes"] if ext in AUDIO | VIDEO else self.settings["max_document_bytes"]
            if size is not None and size > limit:
                raise IngestError("archive_member_size_limit", "Member exceeds expanded byte limit")
            if compressed is not None and size is not None and size / max(1, compressed) > self.settings["archive_ratio"]:
                raise IngestError("archive_ratio_limit", "Member exceeds compression ratio limit")
            output = work / f"member-{ordinal}.{ext.replace('/', '_')}"
            actual = 0
            with opener() as stream, output.open("wb") as dest:
                while block := stream.read(1024 * 1024):
                    actual += len(block)
                    budget["bytes"] += len(block)
                    if actual > limit or budget["bytes"] > self.settings["archive_expanded_bytes"]:
                        raise IngestError("archive_expansion_limit", "Actual expanded bytes exceed resource budget")
                    dest.write(block)
            actual_ext, mime = detect(output, ext)
            member_fields.update(MemberMimeType=mime, DetectedMimeType=mime, SourceMediaType=media_type(actual_ext), FileSizeBytes=actual)
            emitted = False
            for item in self.file(output, actual_ext, member_fields, depth + 1, budget):
                emitted |= isinstance(item, Unit)
                yield item
            if not emitted:
                yield Notice("empty_archive_member", "Archive member produced no usable units", member_fields, "member_skipped")
        except Exception as error:
            yield Notice(getattr(error, "code", "archive_member_failed"), str(error) if isinstance(error, IngestError) else type(error).__name__, member_fields, "member_skipped")
        finally:
            if "output" in locals():
                output.unlink(missing_ok=True)

