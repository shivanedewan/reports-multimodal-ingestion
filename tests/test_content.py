import csv
import io
import json
import zipfile

from PIL import Image

from reports_ingest.chunking import Chunker
from reports_ingest.common import usable_text, extension
from reports_ingest.extraction import Extractor, Notice, Unit
from reports_ingest.events import Events
from reports_ingest.pipeline import source_metadata, iso_date
from .conftest import hit


def units(items):
    return [v for v in items if isinstance(v, Unit)]


def test_missing_with_text_falls_back_as_text(config):
    record = hit(path="missing.mp3", FileExtension="mp3")
    values = list(Extractor(config).record(record))
    unit = units(values)[0]
    assert unit.modality == "text" and unit.fields["SourceMediaType"] == "audio"
    assert unit.fields["UsedTextFallback"] and unit.fields["FileAvailability"] == "missing"
    assert not unit.fields["PreviewAvailable"]


def test_missing_without_text_skips(config):
    values = list(Extractor(config).record(hit(text="  ", path="gone.zip")))
    assert not units(values)
    assert values[0].code == "missing_file_no_text"


def test_short_useful_text_preserved():
    assert usable_text("X") and usable_text("1234")
    assert not usable_text("\x00" * 50 + "x")
    assert not usable_text("Text extraction failed")


def test_token_chunks_cover_every_character_and_obey_limit(config, tokenizer):
    text = "\n\n".join("word " * 300 + str(i) + "." for i in range(20))
    chunks = list(Chunker(tokenizer, config.chunk).chunks(text, "Very long title " * 200))
    covered = set()
    for item in chunks:
        assert item.tokens <= 1400
        assert item.text == text[item.start:item.end]
        covered.update(range(item.start, item.end))
    assert covered == set(range(len(text)))
    assert len(chunks) > 5
    assert any(a.end > b.start for a, b in zip(chunks, chunks[1:]))


def test_single_small_document_is_one_chunk(config, tokenizer):
    assert len(list(Chunker(tokenizer, config.chunk).chunks("12345", "none"))) == 1


def test_zip_valid_bad_and_unsafe_members_are_independent(config, tmp_path):
    archive = tmp_path / "legacy.zip"
    with zipfile.ZipFile(archive, "w") as output:
        output.writestr("good.txt", "Useful legacy document")
        output.writestr("broken.pdf", b"%PDF not a real pdf")
        output.writestr("../outside.txt", "Never extract this")
        output.writestr("sub/good.txt", "Another member")
    values = list(Extractor(config).record(hit(path=str(archive), text="")))
    assert len(units(values)) == 2
    assert len({v.fields["MemberId"] for v in units(values)}) == 2
    assert any(isinstance(v, Notice) and v.code == "unsafe_archive_member" for v in values)
    assert not (tmp_path.parent / "outside.txt").exists()


def test_zip_actual_expansion_limit_does_not_stop_next_source(config, tmp_path):
    config.extract["archive_expanded_bytes"] = 10
    archive = tmp_path / "big.zip"
    with zipfile.ZipFile(archive, "w") as output:
        output.writestr("large.txt", "word " * 100)
    values = list(Extractor(config).record(hit(path=str(archive), text="")))
    assert any(isinstance(v, Notice) and v.code == "archive_expansion_limit" for v in values)
    assert units(list(Extractor(config).record(hit("2"))))


def test_odt_package_is_not_generic_archive(config, tmp_path):
    path = tmp_path / "old.odt"
    with zipfile.ZipFile(path, "w") as output:
        output.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        output.writestr("content.xml", '<root><p>Historical evidence</p></root>')
    values = units(list(Extractor(config).record(hit(path=str(path), text=""))))
    assert values[0].text == "Historical evidence"
    assert not values[0].fields.get("MemberId")


def test_docx_native_text(config, tmp_path):
    from docx import Document
    path = tmp_path / "report.docx"
    document = Document(); document.add_paragraph("Evidence in modern Word."); document.save(path)
    values = units(list(Extractor(config).record(hit(path=str(path), text=""))))
    assert "Evidence" in values[0].text


def test_legacy_doc_failure_uses_es_text(config, tmp_path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"unreadable legacy binary")
    # Usable indexed prose is deliberately preferred for ordinary documents.
    values = units(list(Extractor(config).record(hit(path=str(path)))))
    assert values[0].fields["ContentSource"] == "es_text"


def test_image_and_its_associated_text_are_distinct(config, tmp_path):
    path = tmp_path / "scan.png"
    Image.new("RGB", (80, 60), "white").save(path)
    values = units(list(Extractor(config).record(hit(path=str(path), text="OCR text from a report"))))
    assert [v.modality for v in values] == ["image", "text"]
    assert values[0].media.startswith(b"\x89PNG")


def test_spreadsheet_rows_have_headers_and_coordinates(config, tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "table.xlsx"
    book = Workbook(); book.active.append(["Name", "Count"]); book.active.append(["Border patrol", 7]); book.save(path)
    values = units(list(Extractor(config).record(hit(path=str(path), text=""))))
    assert "Name: Border patrol" in values[1].text
    assert values[1].fields["RowStart"] == 2


def test_audio_windows_have_overlap_and_short_final_window(config, tmp_path, monkeypatch):
    import reports_ingest.extraction as module
    extractor = Extractor(config)
    monkeypatch.setattr(extractor, "probe", lambda _: (125, {"audio"}))
    def fake_command(args, timeout):
        duration = float(args[args.index("-t") + 1])
        return b"\x00\x01" * int(duration * 16000)
    monkeypatch.setattr(module, "command", fake_command)
    values = list(extractor.audio(tmp_path / "long.wav", {}))
    assert [(v.fields["StartSeconds"], v.fields["EndSeconds"]) for v in values] == [(0, 60), (50, 110), (100, 125)]
    import wave
    with wave.open(io.BytesIO(values[-1].media), "rb") as wav:
        assert wav.getframerate() == 16000 and wav.getnframes() == 25 * 16000


def test_csv_quotes_doc_id_index_and_reason(config):
    events = Events(config.path(config.runtime["directory"]), "run")
    from reports_ingest.common import IngestError
    events.error({"DocId": "12,34", "SourceIndex": "reports_main_safe"}, IngestError("corrupt", 'Bad "header", newline\nend'), "source_skipped")
    with (events.directory / "ingestion_errors.csv").open(newline="", encoding="utf-8") as file:
        row = next(csv.DictReader(file))
    assert row["DocId"] == "12,34" and row["Reason"].endswith("end")
    assert row["SourceIndex"] == "reports_main_safe"


def test_metadata_keeps_es_id_and_parent_group_and_dates(config):
    record = hit(index="reports_attachment_safe", ParentDocId=12, DocumentDate=1780000000)
    metadata = source_metadata(record, config)
    assert metadata["SourceEsId"] == "es-1" and metadata["RootDocId"] == "12"
    assert metadata["DocumentDate"].startswith("2026-")
    assert iso_date(1780000000000) == iso_date(1780000000)
    assert iso_date("bad-date") is None
    assert "Text" not in metadata["SourceMetadata"]


def test_compound_extension():
    assert extension("old.TAR.GZ") == "tar.gz"


def test_pdf_text_and_scanned_page_anchors(config, tmp_path):
    import pytest
    pymupdf = pytest.importorskip("pymupdf")
    path = tmp_path / "mixed.pdf"
    document = pymupdf.open()
    document.new_page().insert_text((30, 50), "Text evidence on page one")
    document.new_page()
    document.save(path); document.close()
    values = units(list(Extractor(config).record(hit(path=str(path), text=""))))
    assert [(v.modality, v.fields["PageStart"]) for v in values] == [("text", 1), ("image", 2)]


def test_extensionless_utf8_text_is_detected(config, tmp_path):
    path = tmp_path / "README"
    path.write_text("Legacy evidence with no extension", encoding="utf-8")
    values = units(list(Extractor(config).record(hit(path=str(path), text=""))))
    assert values[0].modality == "text" and "Legacy evidence" in values[0].text

